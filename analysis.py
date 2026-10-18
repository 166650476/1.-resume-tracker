"""数据分析与报表导出模块。

功能：
    * 统计投递总量、笔试/面试/Offer 数量与阶段转化率；
    * 使用 matplotlib 生成投递状态饼图和每月投递数量柱状图；
    * 图表自动保存到 export 文件夹；
    * 将全部投递记录导出为格式化的 Excel 报表；
    * 对无数据、异常日期、长期待投递等情况给出友好提示。

统计口径：
    “笔试/面试/Offer 数量”表示当前状态已达到或超过该阶段的记录数；
    “拒绝/落选”记录无法判断具体止步阶段，只计入投递总量与已投递量。
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")  # 无图形界面环境也能生成图片

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import font_manager
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles import Font as ExcelFont
from openpyxl.utils import get_column_letter

import config
import db
import logger


class AnalysisError(RuntimeError):
    """数据分析或报表导出失败。"""


_FONT_CONFIGURED = False


def _configure_chart_font() -> str:
    """为 matplotlib 选择可用的中文字体，避免图表中文显示为方框。"""
    global _FONT_CONFIGURED
    chosen = "DejaVu Sans"
    for font_name in config.CHART_FONT_CANDIDATES:
        try:
            font_manager.findfont(font_name, fallback_to_default=False)
            chosen = font_name
            break
        except (ValueError, RuntimeError):
            continue

    if not _FONT_CONFIGURED:
        plt.rcParams["font.sans-serif"] = [chosen, "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        _FONT_CONFIGURED = True
        if chosen == "DejaVu Sans":
            logger.log_warning("未找到中文字体，图表中文可能显示异常。")
        else:
            logger.log_info("图表中文字体：%s", chosen)
    return chosen


def _safe_rate(numerator: int, denominator: int) -> float:
    """计算百分比并保留两位小数；分母为 0 时返回 0。"""
    if denominator <= 0:
        return 0.0
    return round(numerator / denominator * 100, 2)


def get_statistics() -> dict[str, Any]:
    """统计投递阶段和转化率。

    Returns:
        包含总量、各阶段数量、状态分布和百分比转化率的字典。
        没有记录时返回全 0 结构，不抛出异常。
    """
    records = db.list_applications()
    total = len(records)
    status_counts_all = {status: 0 for status in config.STATUS_OPTIONS}
    raw_counts = Counter(record.get("status", "") for record in records)
    for status, count in raw_counts.items():
        if status in status_counts_all:
            status_counts_all[status] = count

    pending = status_counts_all.get(config.STATUS_PENDING, 0)
    submitted = total - pending
    written = sum(
        status_counts_all.get(status, 0)
        for status in config.WRITTEN_REACHED_STATUSES
    )
    interview = sum(
        status_counts_all.get(status, 0)
        for status in config.INTERVIEW_REACHED_STATUSES
    )
    offer = status_counts_all.get(config.STATUS_OFFER, 0)
    rejected = sum(
        status_counts_all.get(status, 0) for status in config.REJECTED_STATUSES
    )

    stats = {
        "total": total,
        "pending": pending,
        "submitted": submitted,
        "written": written,
        "interview": interview,
        "offer": offer,
        "rejected": rejected,
        "status_counts": status_counts_all,
        "conversion_rates": {
            "简历投递率": _safe_rate(submitted, total),
            "笔试到达率": _safe_rate(written, submitted),
            "面试到达率": _safe_rate(interview, submitted),
            "Offer获得率": _safe_rate(offer, submitted),
            "笔试→面试": _safe_rate(interview, written),
            "面试→Offer": _safe_rate(offer, interview),
        },
    }
    logger.log_info(
        "统计投递数据 | 总数=%s | 笔试=%s | 面试=%s | Offer=%s",
        total,
        written,
        interview,
        offer,
    )
    return stats


def _parse_date(text: Any) -> date | None:
    """安全解析数据库中的日期文本。"""
    try:
        return datetime.strptime(str(text), config.DATE_FORMAT).date()
    except (TypeError, ValueError):
        return None


def detect_data_issues(
    records: list[dict[str, Any]] | None = None,
) -> list[str]:
    """检查数据异常并返回中文提示列表。

    检查项包括：无数据、非法状态、非法日期、非待投递状态下的未来日期、
    超过 30 天未处理的待投递记录、全部记录都停留在待投递阶段。
    """
    records = db.list_applications() if records is None else records
    if not records:
        return ["当前没有投递记录，请先新增记录后再查看统计数据。"]

    issues: list[str] = []
    today = date.today()
    invalid_status_count = 0
    invalid_date_count = 0
    future_count = 0
    stale_pending_count = 0

    for record in records:
        status = record.get("status", "")
        if status not in config.STATUS_OPTIONS:
            invalid_status_count += 1
        parsed = _parse_date(record.get("apply_date", ""))
        if parsed is None:
            invalid_date_count += 1
            continue
        if parsed > today and status != config.STATUS_PENDING:
            future_count += 1
        if (
            status == config.STATUS_PENDING
            and parsed < today - timedelta(days=30)
        ):
            stale_pending_count += 1

    if invalid_status_count:
        issues.append(f"发现 {invalid_status_count} 条无法识别的投递状态，请检查数据。")
    if invalid_date_count:
        issues.append(f"发现 {invalid_date_count} 条无效投递日期，相关记录未计入月度图。")
    if future_count:
        issues.append(
            f"发现 {future_count} 条投递日期在未来的记录，且状态不是“待投递”，请确认。"
        )
    if stale_pending_count:
        issues.append(
            f"有 {stale_pending_count} 条“待投递”记录已超过 30 天，建议及时更新状态。"
        )
    if records and all(
        record.get("status") == config.STATUS_PENDING for record in records
    ):
        issues.append("所有记录都停留在“待投递”，投递转化率暂时没有参考意义。")
    return issues


# ---------------------------------------------------------------------------
# 图表生成
# ---------------------------------------------------------------------------
def _next_export_path(directory: Path, prefix: str, suffix: str) -> Path:
    """生成不覆盖旧报表的导出路径。"""
    stamp = datetime.now().strftime(config.FILE_TIMESTAMP_FORMAT)
    candidate = directory / f"{prefix}{stamp}{suffix}"
    serial = 1
    while candidate.exists():
        candidate = directory / f"{prefix}{stamp}_{serial}{suffix}"
        serial += 1
    return candidate


def _save_status_pie(
    status_counts: dict[str, int],
    output_path: Path,
) -> Path | None:
    """绘制投递状态分布饼图。"""
    items = [
        (status, status_counts.get(status, 0))
        for status in config.STATUS_OPTIONS
        if status_counts.get(status, 0) > 0
    ]
    if not items:
        return None

    labels = [item[0] for item in items]
    values = [item[1] for item in items]
    colors = [config.STATUS_COLORS.get(label, "#90A4AE") for label in labels]

    figure = None
    try:
        figure, axis = plt.subplots(figsize=config.CHART_FIGURE_SIZE)
        wedges, _texts, autotexts = axis.pie(
            values,
            labels=labels,
            colors=colors,
            autopct=lambda pct: f"{pct:.1f}%",
            startangle=90,
            counterclock=False,
            pctdistance=0.72,
            wedgeprops={"linewidth": 1, "edgecolor": "white"},
        )
        for autotext in autotexts:
            autotext.set_fontsize(10)
            autotext.set_color("white")
            autotext.set_fontweight("bold")
        axis.axis("equal")
        axis.set_title(
            "投递状态分布",
            fontsize=16,
            fontweight="bold",
            pad=16,
        )
        axis.legend(
            wedges,
            [f"{label}：{count} 条" for label, count in items],
            title="状态统计",
            loc="center left",
            bbox_to_anchor=(1.02, 0.5),
            frameon=False,
        )
        figure.tight_layout()
        figure.savefig(
            output_path,
            dpi=config.CHART_DPI,
            bbox_inches="tight",
            facecolor="white",
        )
    finally:
        if figure is not None:
            plt.close(figure)
    return output_path


def _save_monthly_bar(
    records: list[dict[str, Any]],
    output_path: Path,
) -> Path | None:
    """绘制每月投递数量柱状图。"""
    data = pd.DataFrame(
        {"apply_date": [record.get("apply_date", "") for record in records]}
    )
    parsed = pd.to_datetime(data["apply_date"], errors="coerce")
    valid = parsed.dropna()
    if valid.empty:
        logger.log_warning("没有有效投递日期，跳过月度柱状图。")
        return None

    counts = valid.dt.to_period("M").astype(str).value_counts().sort_index()
    months = [str(month) for month in counts.index]
    values = [int(value) for value in counts.values]

    figure = None
    try:
        figure, axis = plt.subplots(figsize=config.CHART_FIGURE_SIZE)
        bars = axis.bar(months, values, color="#4C78A8", width=0.62)
        axis.set_title(
            "每月投递数量",
            fontsize=16,
            fontweight="bold",
            pad=16,
        )
        axis.set_xlabel("月份")
        axis.set_ylabel("投递记录数")
        axis.grid(axis="y", linestyle="--", alpha=0.35)
        axis.set_axisbelow(True)
        axis.tick_params(axis="x", rotation=35)
        for bar, value in zip(bars, values):
            axis.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 0.05,
                str(value),
                ha="center",
                va="bottom",
                fontsize=10,
            )
        axis.set_ylim(0, max(values) + 1.2)
        figure.tight_layout()
        figure.savefig(
            output_path,
            dpi=config.CHART_DPI,
            bbox_inches="tight",
            facecolor="white",
        )
    finally:
        if figure is not None:
            plt.close(figure)
    return output_path


def generate_charts(
    output_dir: str | Path | None = None,
) -> dict[str, Path]:
    """生成状态饼图和月度柱状图并保存到 export 目录。

    Returns:
        成功生成的图表路径字典，可能包含 ``status_chart`` 和 ``monthly_chart``。
        无投递记录时返回空字典。
    """
    records = db.list_applications()
    if not records:
        logger.log_warning("生成图表时没有投递记录")
        return {}

    target_directory = (
        Path(output_dir).resolve() if output_dir else config.EXPORT_DIR
    )

    try:
        target_directory.mkdir(parents=True, exist_ok=True)
        _configure_chart_font()
        stats = get_statistics()
        results: dict[str, Path] = {}

        status_path = _next_export_path(
            target_directory, config.STATUS_CHART_PREFIX, ".png"
        )
        saved_status = _save_status_pie(stats["status_counts"], status_path)
        if saved_status:
            results["status_chart"] = saved_status

        monthly_path = _next_export_path(
            target_directory, config.MONTHLY_CHART_PREFIX, ".png"
        )
        saved_monthly = _save_monthly_bar(records, monthly_path)
        if saved_monthly:
            results["monthly_chart"] = saved_monthly
    except (OSError, ValueError, RuntimeError) as exc:
        logger.log_exception("生成统计图表", exc)
        raise AnalysisError(f"生成统计图表失败：{exc}") from exc

    logger.log_operation(
        "生成统计图表",
        " | ".join(f"{key}={value}" for key, value in results.items()),
    )
    return results


# ---------------------------------------------------------------------------
# Excel 报表导出
# ---------------------------------------------------------------------------
def _build_application_dataframe(
    records: list[dict[str, Any]],
) -> pd.DataFrame:
    """把投递记录转换成中文表头的 DataFrame。"""
    rows = []
    for record in records:
        rows.append(
            {
                "ID": record.get("id", ""),
                "公司名称": record.get("company_name", ""),
                "岗位名称": record.get("position", ""),
                "投递日期": record.get("apply_date", ""),
                "投递状态": record.get("status", ""),
                "备注": record.get("notes", ""),
                "关联模板ID": record.get("resume_template_id") or "",
                "关联模板名称": record.get("template_name") or "",
                "创建时间": record.get("created_at", ""),
                "更新时间": record.get("updated_at", ""),
            }
        )
    return pd.DataFrame(rows)


def _build_summary_dataframe(stats: dict[str, Any]) -> pd.DataFrame:
    """构造统计概览工作表。"""
    rows = [
        ("投递记录总数", stats["total"]),
        ("待投递", stats["pending"]),
        ("已投递（不含待投递）", stats["submitted"]),
        ("笔试到达数量", stats["written"]),
        ("面试到达数量", stats["interview"]),
        ("Offer数量", stats["offer"]),
        ("拒绝/落选数量", stats["rejected"]),
    ]
    for label, value in stats["conversion_rates"].items():
        rows.append((label, f"{value:.2f}%"))
    return pd.DataFrame(rows, columns=["指标", "数值"])


def _build_status_dataframe(stats: dict[str, Any]) -> pd.DataFrame:
    """构造状态分布工作表。"""
    rows = [
        {"投递状态": status, "记录数量": count}
        for status, count in stats["status_counts"].items()
    ]
    return pd.DataFrame(rows)


def _excel_text_width(value: Any) -> int:
    """粗略计算单元格显示宽度，让中文列宽更合适。"""
    text = str(value or "")
    return sum(2 if ord(char) > 127 else 1 for char in text)


def _format_worksheet(worksheet: Any) -> None:
    """美化 Excel 工作表：表头底色、冻结首行、自动筛选、列宽自适应。"""
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = ExcelFont(color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24

    for column_index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_width = 10
        for cell in worksheet[letter]:
            display_width = _excel_text_width(cell.value)
            max_width = max(max_width, display_width)
        worksheet.column_dimensions[letter].width = min(max_width + 2, 48)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment


def export_applications_excel(
    output_dir: str | Path | None = None,
) -> Path | None:
    """将全部投递记录导出为 Excel 文件。

    Excel 包含三个工作表：
        * 投递记录：全部明细；
        * 统计概览：数量及阶段转化率；
        * 状态分布：各状态数量。

    Args:
        output_dir: 输出目录，默认 export 文件夹。

    Returns:
        导出文件路径；无数据时返回 None。
    """
    records = db.list_applications()
    if not records:
        logger.log_warning("导出 Excel 时没有投递记录")
        return None

    target_directory = (
        Path(output_dir).resolve() if output_dir else config.EXPORT_DIR
    )
    output_path: Path | None = None
    try:
        target_directory.mkdir(parents=True, exist_ok=True)
        output_path = _next_export_path(
            target_directory, config.EXCEL_REPORT_PREFIX, ".xlsx"
        )
        stats = get_statistics()
        detail_frame = _build_application_dataframe(records)
        summary_frame = _build_summary_dataframe(stats)
        status_frame = _build_status_dataframe(stats)

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            detail_frame.to_excel(writer, sheet_name="投递记录", index=False)
            summary_frame.to_excel(writer, sheet_name="统计概览", index=False)
            status_frame.to_excel(writer, sheet_name="状态分布", index=False)

        workbook = load_workbook(output_path)
        try:
            for worksheet in workbook.worksheets:
                _format_worksheet(worksheet)
            workbook.save(output_path)
        finally:
            workbook.close()
    except (OSError, PermissionError, ValueError, KeyError, ImportError) as exc:
        logger.log_exception("导出 Excel 投递报表", exc)
        if output_path and output_path.exists():
            try:
                output_path.unlink()
            except OSError:
                pass
        raise AnalysisError(f"导出 Excel 投递报表失败：{exc}") from exc
    except Exception as exc:
        logger.log_exception("导出 Excel 投递报表", exc)
        if output_path and output_path.exists():
            try:
                output_path.unlink()
            except OSError:
                pass
        raise AnalysisError(f"Excel 报表生成异常：{exc}") from exc

    logger.log_operation(
        "导出 Excel 投递报表",
        f"记录数={len(records)} | 输出={output_path}",
    )
    return output_path
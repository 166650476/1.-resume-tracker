"""项目自动化全流程自测脚本。

运行方式（在项目根目录）：
    python selftest.py

脚本使用 export 文件夹中的临时数据库，不会修改正式 resume.db。
覆盖流程：
    个人信息录入 → 模板新建/复制/修改/预览 → 单份与批量 Word 生成 →
    投递新增/重复检测/状态更新 → 批量导入 → 统计图表/Excel/txt →
    数据库备份与恢复 → 清理测试数据。
"""

from __future__ import annotations

import sys
import traceback
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document as WordDocument

import analysis
import config
import db
import logger
import resume_generator
import tracker

TEST_DB_PATH = config.EXPORT_DIR / "_selftest_resume.db"
TEST_CSV_PATH = config.EXPORT_DIR / "_selftest_import.csv"
ARTIFACTS: set[Path] = set()
PASSED_STEPS: list[str] = []


def _check(condition: bool, message: str) -> None:
    """断言辅助函数，失败时抛出带中文说明的异常。"""
    if not condition:
        raise RuntimeError(message)


def _step(title: str) -> None:
    """输出自测步骤标题。"""
    print(f"\n[自测] {title}")


def _pass(message: str) -> None:
    """记录并输出通过项。"""
    PASSED_STEPS.append(message)
    print(f"  [通过] {message}")


def _remember(path: str | Path | None) -> None:
    """记录需要清理的测试产物。"""
    if path:
        ARTIFACTS.add(Path(path).resolve())


def _remove_file(path: Path) -> None:
    """安静删除单个测试文件。"""
    try:
        if path.exists() and path.is_file():
            path.unlink()
    except OSError:
        pass


def _cleanup_test_database() -> None:
    """删除临时数据库及其可能的日志文件。"""
    for suffix in ("", "-wal", "-shm", "-journal"):
        _remove_file(Path(f"{TEST_DB_PATH}{suffix}"))


def _cleanup() -> None:
    """清理本次自测产生的全部文件。"""
    for path in ARTIFACTS:
        _remove_file(path)
    _remove_file(TEST_CSV_PATH)
    _cleanup_test_database()


def _event_date(days_ago: int) -> str:
    """返回相对今天偏移的日期文本。"""
    return (date.today() - timedelta(days=days_ago)).strftime(config.DATE_FORMAT)


def run_selftest() -> None:
    """执行完整业务流程自测。"""
    config.ensure_directories()
    _cleanup_test_database()
    db.set_database_path(TEST_DB_PATH)
    created = db.init_database()
    _check(created, "临时数据库未正确创建。")
    _pass("自动创建临时测试数据库")

    # ------------------------------------------------------------------
    _step("1. 录入个人基础信息与输入校验")
    person_payload = {
        "name": "测试同学",
        "phone": "13900000001",
        "email": "test.resume@example.com",
        "gender": "男",
        "birth_date": "2002-05-20",
        "city": "杭州",
        "job_intention": "Python后端开发",
        "education": "本科",
        "school": "测试大学",
        "major": "软件工程",
        "graduation_year": "2026",
    }
    person_id = db.create_personal_info(person_payload)
    _check(person_id > 0, "个人信息未写入。")
    _check(
        db.get_personal_info(person_id)["name"] == "测试同学",
        "个人信息读取不一致。",
    )
    _pass("录入并读取个人基础信息")

    try:
        db.create_personal_info(person_payload)
    except db.DuplicateRecordError:
        _pass("手机号重复检测生效")
    else:
        raise RuntimeError("重复手机号未被拦截。")

    try:
        invalid = dict(person_payload)
        invalid["phone"] = "12345"
        invalid["email"] = "bad-email"
        db.create_personal_info(invalid)
    except db.ValidationError:
        _pass("手机号/邮箱格式校验生效")
    else:
        raise RuntimeError("非法手机号或邮箱未被拦截。")

    # ------------------------------------------------------------------
    _step("2. 新建、复制并修改简历模板")
    template_payload: dict[str, Any] = {
        "template_name": "Python开发工程师模板",
        "personal_info_id": person_id,
        "target_position": "Python开发工程师",
        "education_experience": (
            "2022.09-2026.06  测试大学  软件工程  本科\n"
            "- 主修课程：数据结构、数据库原理、操作系统"
        ),
        "internship_experience": (
            "2025.07-2025.09  示例科技有限公司  后端实习生\n"
            "- 使用 Python 开发数据接口\n"
            "- 参与 SQLite 数据建模与查询优化"
        ),
        "project_experience": (
            "2025.03-2025.06  简历批量生成与投递管理系统\n"
            "- 使用 sqlite3 设计三张数据表\n"
            "- 使用 python-docx 生成规范 Word 简历"
        ),
        "skills": (
            "- 熟悉 Python、SQLite、pandas、matplotlib\n"
            "- 熟悉 Git 与 Linux 基础命令\n"
            "- 具备良好的文档编写习惯"
        ),
        "awards": "- 校级一等奖学金（2024）\n- 全国大学生软件竞赛省级二等奖",
        "self_evaluation": "学习能力强，沟通积极，能够快速适应团队协作与项目节奏。",
        "font_name": "微软雅黑",
        "font_size": 10.5,
        "line_spacing": 1.25,
        "paragraph_indent": 0.74,
    }
    template_id = db.create_resume_template(template_payload)
    _check(template_id > 0, "简历模板未写入。")
    _pass("新建简历模板")

    try:
        db.create_resume_template(template_payload)
    except db.DuplicateRecordError:
        _pass("模板名称重复检测生效")
    else:
        raise RuntimeError("重复模板名称未被拦截。")

    copied_id, copied_name = resume_generator.copy_template(template_id)
    copied_template = db.get_resume_template(copied_id)
    _check(copied_template is not None, "复制后的模板不存在。")
    _check(
        copied_template["skills"] == template_payload["skills"],
        "模板复制内容不完整。",
    )
    _check(
        copied_name.endswith(config.TEMPLATE_COPY_SUFFIX),
        "副本名称不符合预期。",
    )
    _pass("复制模板并保持全部内容")

    db.update_resume_template(
        copied_id,
        {
            "target_position": "Python后端开发工程师",
            "self_evaluation": "具备扎实的 Python 基础，乐于沟通，关注代码质量与交付效率。",
        },
    )
    updated_template = db.get_resume_template(copied_id)
    _check(
        updated_template is not None
        and updated_template["target_position"] == "Python后端开发工程师",
        "模板修改未生效。",
    )
    _pass("修改复制后的模板内容")

    # ------------------------------------------------------------------
    _step("3. 预览简历并生成 Word")
    preview_text = resume_generator.render_preview_text(copied_id)
    _check("Python后端开发工程师" in preview_text, "预览缺少目标岗位。")
    _check("【项目经历】" in preview_text, "预览缺少项目经历区块。")
    _check("• 熟悉 Python" in preview_text, "预览未显示项目符号。")
    _pass("命令行文本预览")

    doc_path = resume_generator.generate_resume(copied_id)
    _remember(doc_path)
    _check(doc_path.exists() and doc_path.stat().st_size > 0, "Word 简历未生成。")
    generated_doc = WordDocument(str(doc_path))
    doc_text = "\n".join(paragraph.text for paragraph in generated_doc.paragraphs)
    _check("测试同学" in doc_text, "Word 简历缺少姓名。")
    _check("Python后端开发工程师" in doc_text, "Word 简历缺少目标岗位。")
    _check("项目经历" in doc_text, "Word 简历缺少项目经历。")
    has_bullet = any(
        paragraph.style.name == "List Bullet"
        for paragraph in generated_doc.paragraphs
    )
    _check(has_bullet, "Word 简历未生成项目符号段落。")
    _pass("生成 Word 简历并校验内容/项目符号")

    batch_results = resume_generator.batch_generate_all()
    _check(len(batch_results) == 2, "批量生成未覆盖全部两个模板。")
    for result in batch_results:
        _remember(result["path"])
        _check(
            result["path"] is not None,
            f"模板 {result['template_name']} 批量生成失败。",
        )
        _check(
            result["path"].exists(),
            f"模板 {result['template_name']} 文件不存在。",
        )
    _pass("批量生成全部模板 Word 文档")

    # ------------------------------------------------------------------
    _step("4. 新增投递记录、重复检测与状态更新")
    first_app_id = tracker.add_application(
        company_name="示例科技有限公司",
        position="Python后端开发工程师",
        apply_date=_event_date(0),
        status="待投递",
        notes="通过校园招聘官网投递。",
        resume_template_id=copied_id,
    )
    try:
        tracker.add_application(
            company_name="示例科技有限公司",
            position="Python后端开发工程师",
            apply_date=_event_date(0),
            status="待投递",
            notes="重复记录",
            resume_template_id=copied_id,
        )
    except db.DuplicateRecordError:
        _pass("投递记录重复检测生效")
    else:
        raise RuntimeError("重复投递记录未被拦截。")

    tracker.update_status(
        first_app_id,
        "面试",
        "一面问题：SQLite 索引与事务；薪资沟通：期望 9K-11K。",
    )
    first_app = db.get_application(first_app_id)
    _check(first_app is not None and first_app["status"] == "面试", "投递状态未更新。")
    _check("一面问题" in first_app["notes"], "面试备注未追加。")
    _pass("新增投递并更新状态、追加面试/薪资备注")

    second_app_id = tracker.add_application(
        company_name="未来数据科技有限公司",
        position="数据分析师",
        apply_date=_event_date(32),
        status="收到offer",
        notes="已收到书面 Offer，正在评估薪资与通勤。",
        resume_template_id=copied_id,
    )
    _check(second_app_id > 0, "第二条投递记录未新增。")
    _pass("新增第二条不同状态/月份的投递记录")

    # ------------------------------------------------------------------
    _step("5. 批量导入投递记录（含重复行与非法行）")
    csv_content = (
        "公司名称,岗位,投递日期,投递状态,备注,简历模板ID\n"
        f"批量导入公司,数据分析师,{_event_date(0)},已投递,批量导入有效记录,{copied_id}\n"
        f"示例科技有限公司,Python后端开发工程师,{_event_date(0)},待投递,重复记录,{copied_id}\n"
        f"校验失败公司,测试岗位,{_event_date(0)},未知状态,非法状态,\n"
    )
    TEST_CSV_PATH.write_text(csv_content, encoding="utf-8-sig")
    _remember(TEST_CSV_PATH)
    import_result = tracker.batch_import_applications(TEST_CSV_PATH)
    _check(import_result["success_count"] == 1, "批量导入成功数不正确。")
    _check(import_result["duplicate_count"] == 1, "批量导入重复检测数量不正确。")
    _check(import_result["failed_count"] == 1, "批量导入非法行统计不正确。")
    _pass("批量导入：成功 1 条、重复 1 条、失败 1 条")

    # ------------------------------------------------------------------
    _step("6. 统计分析、图表、Excel 与 txt 导出")
    stats = analysis.get_statistics()
    _check(stats["total"] == 3, "投递总量统计不正确。")
    _check(stats["submitted"] == 3, "已投递数量统计不正确。")
    _check(stats["written"] == 2, "笔试到达数量统计不正确。")
    _check(stats["interview"] == 2, "面试到达数量统计不正确。")
    _check(stats["offer"] == 1, "Offer 数量统计不正确。")
    _check(stats["conversion_rates"]["Offer获得率"] > 0, "Offer 转化率计算异常。")
    _pass("统计投递总量、阶段数量与转化率")

    charts = analysis.generate_charts()
    _check("status_chart" in charts, "未生成状态饼图。")
    _check("monthly_chart" in charts, "未生成月度柱状图。")
    for chart_path in charts.values():
        _remember(chart_path)
        _check(
            chart_path.exists() and chart_path.stat().st_size > 0,
            "图表文件为空。",
        )
    _pass("生成状态饼图和每月投递柱状图")

    excel_path = analysis.export_applications_excel()
    _remember(excel_path)
    _check(excel_path is not None and excel_path.exists(), "Excel 报表未生成。")
    excel_sheets = pd.read_excel(excel_path, sheet_name=None)
    _check(
        set(excel_sheets) == {"投递记录", "统计概览", "状态分布"},
        "Excel 工作表不完整。",
    )
    _check(len(excel_sheets["投递记录"]) == 3, "Excel 明细行数不正确。")
    _pass("导出 Excel 投递报表（3 个工作表）")

    txt_path = tracker.export_applications_txt()
    _remember(txt_path)
    _check(txt_path is not None and txt_path.exists(), "txt 备份未生成。")
    txt_content = txt_path.read_text(encoding="utf-8-sig")
    _check("示例科技有限公司" in txt_content, "txt 备份缺少投递记录。")
    _pass("导出投递记录 txt 备份")

    # ------------------------------------------------------------------
    _step("7. 数据库备份与恢复")
    backup_path = db.backup_database()
    _remember(backup_path)
    _check(
        backup_path.exists() and backup_path.stat().st_size > 0,
        "数据库备份未生成。",
    )
    _pass("在线备份数据库")

    db.update_application(second_app_id, {"status": "落选"})
    _check(
        db.get_application(second_app_id)["status"] == "落选",
        "恢复测试前的状态修改失败。",
    )
    safety_backup = db.restore_database(backup_path)
    _remember(safety_backup)
    restored = db.get_application(second_app_id)
    _check(
        restored is not None and restored["status"] == "收到offer",
        "数据库恢复结果不正确。",
    )
    _check(safety_backup.exists(), "恢复前安全备份未生成。")
    _pass("从备份恢复数据库，并验证数据回到备份时状态")

    counts = db.get_record_counts()
    _check(
        counts
        == {
            "personal_info": 1,
            "resume_templates": 2,
            "applications": 3,
        },
        "恢复后记录数异常。",
    )
    _pass("恢复后三张业务表记录数一致")


def main() -> int:
    """运行自测并返回退出码。"""
    print(config.MENU_SEPARATOR)
    print("  简历批量生成与投递管理系统 - 全流程自测")
    print(config.MENU_SEPARATOR)
    logger.log_info("开始执行全流程自测")
    success = False
    try:
        run_selftest()
        success = True
    except Exception as exc:
        logger.log_exception("全流程自测", exc)
        print(f"\n[自测失败] {exc}")
        traceback.print_exc()
    finally:
        db.set_database_path(None)
        _cleanup()
        db.init_database()

    if success:
        print("\n" + config.MENU_SEPARATOR)
        print(f"  [全部通过] 共完成 {len(PASSED_STEPS)} 项功能验证")
        print("  测试数据与临时文件已清理，正式数据库未写入测试记录。")
        print(config.MENU_SEPARATOR)
        logger.log_info("全流程自测通过，共 %s 项", len(PASSED_STEPS))
        return 0

    print("\n[自测结果] 存在失败项，请查看 logs 目录中的日志。")
    logger.log_error("全流程自测失败")
    return 1


if __name__ == "__main__":
    sys.exit(main())
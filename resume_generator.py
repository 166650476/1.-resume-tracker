"""简历生成与模板管理模块。

主要能力：
    * 基于数据库中的简历模板与个人资料生成规范 Word 简历；
    * 支持自定义字体、字号、行距、首行缩进；
    * 支持以 ``-``、``*``、``•`` 开头的行自动转换为项目符号；
    * 在命令行直接预览简历文本；
    * 复制已有模板并自动生成不重复的副本名称；
    * 一键批量生成全部模板对应的 Word 文档。

生成的 docx 文件默认输出到项目根目录的 resume_output 文件夹。
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config
import db
import logger


class ResumeGenerationError(RuntimeError):
    """简历生成失败。"""


BULLET_PREFIXES = ("-", "*", "•", "·")


def _safe_filename(text: str, max_length: int = 120) -> str:
    """移除 Windows 文件名中不允许的字符，并限制文件名长度。"""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    return (cleaned[:max_length] or "resume").rstrip(" .")


def _get_template_and_person(template_id: int) -> tuple[dict[str, Any], dict[str, Any]]:
    """读取模板及其关联的个人信息，缺失时抛出中文异常。"""
    template = db.get_resume_template(template_id)
    if template is None:
        raise LookupError(f"ID 为 {template_id} 的简历模板不存在。")
    person = db.get_personal_info(int(template["personal_info_id"]))
    if person is None:
        raise LookupError("简历模板关联的个人基础信息不存在，请先补充个人信息。")
    return template, person


def _set_run_font(
    run: Any,
    font_name: str,
    font_size: float | None = None,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    """为 Word run 同时设置西文和东亚字体，避免中文回退为宋体。"""
    run.font.name = font_name
    if font_size is not None:
        run.font.size = Pt(font_size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # python-docx 默认只设置 w:ascii，中文还需要 w:eastAsia
    if run._element.rPr is not None and run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _apply_paragraph_format(
    paragraph: Any,
    line_spacing: float,
    first_line_indent_cm: float = 0.0,
    space_before: float = 0.0,
    space_after: float = 3.0,
    alignment: Any | None = None,
) -> None:
    """统一设置段落格式。"""
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing = line_spacing
    paragraph_format.space_before = Pt(space_before)
    paragraph_format.space_after = Pt(space_after)
    if first_line_indent_cm > 0:
        paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    if alignment is not None:
        paragraph.alignment = alignment


def _configure_document(document: Any, font_name: str, font_size: float) -> None:
    """设置页面边距和正文默认字体。"""
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal_style = document.styles["Normal"]
    normal_style.font.name = font_name
    normal_style.font.size = Pt(font_size)
    if (
        normal_style.element.rPr is not None
        and normal_style.element.rPr.rFonts is not None
    ):
        normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _add_section_heading(
    document: Any,
    title: str,
    font_name: str,
    font_size: float,
) -> None:
    """添加带主题色和粗体效果的区块标题。"""
    paragraph = document.add_paragraph()
    _apply_paragraph_format(
        paragraph,
        line_spacing=1.0,
        space_before=8,
        space_after=4,
    )
    run = paragraph.add_run(f"▌ {title}")
    _set_run_font(
        run,
        font_name,
        font_size + 1.5,
        bold=True,
        color=RGBColor(31, 78, 121),
    )


def _add_text_block(
    document: Any,
    content: str,
    font_name: str,
    font_size: float,
    line_spacing: float,
    paragraph_indent: float,
) -> None:
    """把多行文本写入 Word，并识别项目符号行。

    Examples:
        普通行会作为正文段落；以 ``-`` 开头的行会生成 Word 项目符号。
    """
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith(BULLET_PREFIXES):
            bullet_text = line[1:].strip()
            if not bullet_text:
                continue
            paragraph = document.add_paragraph(style="List Bullet")
            _apply_paragraph_format(
                paragraph,
                line_spacing=line_spacing,
                space_after=2,
            )
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.37)
            run = paragraph.add_run(bullet_text)
        else:
            paragraph = document.add_paragraph()
            _apply_paragraph_format(
                paragraph,
                line_spacing=line_spacing,
                first_line_indent_cm=paragraph_indent,
                space_after=3,
            )
            run = paragraph.add_run(line)
        _set_run_font(run, font_name, font_size)


def _education_fallback(person: dict[str, Any]) -> str:
    """当模板未填写教育背景时，用个人基础信息拼接一条教育经历。"""
    parts: list[str] = []
    if person.get("school"):
        parts.append(person["school"])
    if person.get("major"):
        parts.append(person["major"])
    if person.get("education"):
        parts.append(person["education"])
    if person.get("graduation_year"):
        parts.append(f"{person['graduation_year']} 年毕业")
    return " | ".join(parts)


def _build_sections(
    template: dict[str, Any],
    person: dict[str, Any],
) -> list[tuple[str, str]]:
    """组合简历区块，供 Word 生成和文本预览共同使用。"""
    sections: list[tuple[str, str]] = []

    intention_lines: list[str] = []
    if template.get("target_position"):
        intention_lines.append(f"目标岗位：{template['target_position']}")
    if person.get("job_intention"):
        intention_lines.append(f"求职方向：{person['job_intention']}")
    if person.get("city"):
        intention_lines.append(f"期望城市：{person['city']}")
    if intention_lines:
        sections.append(("求职意向", "\n".join(intention_lines)))

    education_content = (template.get("education_experience") or "").strip()
    if not education_content:
        education_content = _education_fallback(person)
    if education_content:
        sections.append(("教育背景", education_content))

    for field, label in config.RESUME_SECTION_LABELS:
        if field == "education_experience":
            continue
        content = (template.get(field) or "").strip()
        if content:
            sections.append((label, content))
    return sections


def render_preview_text(template_id: int) -> str:
    """生成可直接打印到命令行的纯文本简历预览。

    Args:
        template_id: 简历模板 ID。

    Returns:
        包含个人信息、各区块和排版参数的纯文本。
    """
    template, person = _get_template_and_person(template_id)
    contact_parts = [person.get("phone", ""), person.get("email", "")]
    if person.get("city"):
        contact_parts.append(person["city"])
    contact_text = " | ".join(part for part in contact_parts if part)

    lines: list[str] = [
        person.get("name", ""),
        f"目标岗位：{template.get('target_position', '')}",
        f"联系方式：{contact_text}",
    ]

    for label, content in _build_sections(template, person):
        lines.append("")
        lines.append(f"【{label}】")
        for raw_line in content.splitlines():
            text = raw_line.strip()
            if not text:
                continue
            if text.startswith(BULLET_PREFIXES):
                lines.append("  • " + text[1:].strip())
            else:
                lines.append("  " + text)

    lines.extend(
        [
            "",
            "-" * 60,
            (
                "排版参数："
                f"字体={template.get('font_name')} | "
                f"字号={template.get('font_size')} | "
                f"行距={template.get('line_spacing')} | "
                f"首行缩进={template.get('paragraph_indent')} cm"
            ),
        ]
    )
    logger.log_operation("预览简历模板", f"模板ID={template_id}")
    return "\n".join(lines)


def preview_resume(template_id: int) -> None:
    """在终端打印简历预览文本。"""
    print(config.SUB_SEPARATOR)
    print(render_preview_text(template_id))
    print(config.SUB_SEPARATOR)


def _unique_output_path(output_dir: Path, stem: str) -> Path:
    """生成不覆盖已有文件的 docx 输出路径。"""
    candidate = output_dir / f"{stem}.docx"
    serial = 1
    while candidate.exists():
        candidate = output_dir / f"{stem}_{serial}.docx"
        serial += 1
    return candidate


def generate_resume(
    template_id: int,
    output_dir: str | Path | None = None,
    filename: str | None = None,
) -> Path:
    """生成一份 Word 简历。

    Args:
        template_id: 简历模板 ID。
        output_dir: 输出目录，默认 resume_output。
        filename: 自定义文件名；不传则自动生成包含姓名、岗位和时间戳的文件名。

    Returns:
        生成的 docx 文件绝对路径。

    Raises:
        LookupError: 模板或个人资料不存在。
        ResumeGenerationError: Word 文档创建或保存失败。
    """
    template, person = _get_template_and_person(template_id)
    target_directory = (
        Path(output_dir).resolve() if output_dir else config.RESUME_OUTPUT_DIR
    )

    try:
        target_directory.mkdir(parents=True, exist_ok=True)
        if filename:
            safe_stem = _safe_filename(Path(filename).stem)
        else:
            stamp = datetime.now().strftime(config.FILE_TIMESTAMP_FORMAT)
            safe_stem = _safe_filename(
                f"{person.get('name', '简历')}_{template.get('target_position', '')}_"
                f"{template.get('template_name', '')}_{stamp}"
            )
        output_path = _unique_output_path(target_directory, safe_stem)

        font_name = str(template.get("font_name") or config.DEFAULT_FONT_NAME)
        font_size = float(template.get("font_size") or config.DEFAULT_FONT_SIZE)
        line_spacing = float(
            template.get("line_spacing") or config.DEFAULT_LINE_SPACING
        )
        paragraph_indent = float(template.get("paragraph_indent") or 0.0)

        document = Document()
        _configure_document(document, font_name, font_size)
        document.core_properties.title = f"{person.get('name', '')}的求职简历"
        document.core_properties.author = person.get("name", "")
        document.core_properties.subject = template.get("target_position", "")

        name_paragraph = document.add_paragraph()
        _apply_paragraph_format(
            name_paragraph,
            line_spacing=1.0,
            space_after=2,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        name_run = name_paragraph.add_run(person.get("name", ""))
        _set_run_font(
            name_run,
            font_name,
            font_size + 9,
            bold=True,
            color=RGBColor(31, 78, 121),
        )

        position_paragraph = document.add_paragraph()
        _apply_paragraph_format(
            position_paragraph,
            line_spacing=1.0,
            space_after=2,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        position_run = position_paragraph.add_run(
            template.get("target_position", "")
        )
        _set_run_font(position_run, font_name, font_size + 1.5, bold=True)

        contact_parts = [person.get("phone", ""), person.get("email", "")]
        if person.get("city"):
            contact_parts.append(person["city"])
        contact_paragraph = document.add_paragraph()
        _apply_paragraph_format(
            contact_paragraph,
            line_spacing=1.0,
            space_after=4,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        contact_run = contact_paragraph.add_run(
            " | ".join(part for part in contact_parts if part)
        )
        _set_run_font(contact_run, font_name, font_size - 0.5)

        for label, content in _build_sections(template, person):
            _add_section_heading(document, label, font_name, font_size)
            _add_text_block(
                document,
                content,
                font_name,
                font_size,
                line_spacing,
                paragraph_indent,
            )

        document.save(str(output_path))
    except (OSError, PermissionError, ValueError, KeyError) as exc:
        logger.log_exception("生成 Word 简历", exc)
        raise ResumeGenerationError(f"生成 Word 简历失败：{exc}") from exc
    except Exception as exc:
        logger.log_exception("生成 Word 简历", exc)
        raise ResumeGenerationError(f"Word 文档生成异常：{exc}") from exc

    logger.log_operation(
        "生成 Word 简历",
        f"模板ID={template_id} | 输出={output_path}",
    )
    return output_path


def _make_unique_copy_name(source_name: str) -> str:
    """基于原模板名生成不重复的副本名称。"""
    suffix = config.TEMPLATE_COPY_SUFFIX
    serial = 0
    while True:
        current_suffix = suffix if serial == 0 else f"（副本{serial + 1}）"
        available_length = max(
            1,
            config.TEMPLATE_NAME_MAX_LEN - len(current_suffix),
        )
        candidate = source_name[:available_length].rstrip() + current_suffix
        if not db.template_name_exists(candidate):
            return candidate
        serial += 1


def copy_template(source_template_id: int, new_name: str | None = None) -> tuple[int, str]:
    """复制已有简历模板，快速创建一份内容相同的新模板。

    Args:
        source_template_id: 被复制的模板 ID。
        new_name: 新模板名称；为空时自动使用“原名称（副本）”。

    Returns:
        (新模板 ID, 新模板名称)

    Raises:
        LookupError: 源模板不存在。
        DuplicateRecordError: 手动指定的新名称已存在。
    """
    source = db.get_resume_template(source_template_id)
    if source is None:
        raise LookupError(f"ID 为 {source_template_id} 的简历模板不存在。")

    target_name = (new_name or "").strip()
    if not target_name:
        target_name = _make_unique_copy_name(source["template_name"])

    payload = {
        field: source.get(field)
        for field in db.TEMPLATE_FIELDS
    }
    payload["template_name"] = target_name
    new_template_id = db.create_resume_template(payload)
    logger.log_operation(
        "复制简历模板",
        f"源模板={source_template_id} | 新模板={new_template_id} | 名称={target_name}",
    )
    return new_template_id, target_name


def batch_generate_all(
    output_dir: str | Path | None = None,
) -> list[dict[str, Any]]:
    """为全部简历模板批量生成 Word 文档。

    单个模板失败不会中断其他模板，结果中会分别记录成功路径或错误原因。

    Returns:
        每个模板一条结果记录：
        ``{"template_id", "template_name", "path", "error"}``。
    """
    templates = db.list_resume_templates()
    if not templates:
        logger.log_warning("批量生成简历时没有可用模板")
        return []

    results: list[dict[str, Any]] = []
    for template in templates:
        template_id = int(template["id"])
        template_name = template["template_name"]
        try:
            output_path = generate_resume(template_id, output_dir)
            results.append(
                {
                    "template_id": template_id,
                    "template_name": template_name,
                    "path": output_path,
                    "error": "",
                }
            )
        except Exception as exc:  # 单个失败不阻断整批任务
            logger.log_exception(
                f"批量生成模板“{template_name}”的 Word 简历", exc
            )
            results.append(
                {
                    "template_id": template_id,
                    "template_name": template_name,
                    "path": None,
                    "error": str(exc),
                }
            )

    success_count = sum(1 for item in results if item["path"] is not None)
    logger.log_operation(
        "批量生成简历",
        f"模板总数={len(results)} | 成功={success_count} | 失败={len(results) - success_count}",
    )
    return results